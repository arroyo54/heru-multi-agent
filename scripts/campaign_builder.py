"""
Constructor de Campañas — heru.app
Genera: PPT estratégico + CSVs por canal (Google, Meta, TikTok, YouTube)

Uso:
  python scripts/campaign_builder.py
"""
import os, sys, csv, re, json
from pathlib import Path
from datetime import datetime

sys.path.insert(0, str(Path(__file__).parent.parent))
from dotenv import load_dotenv
if Path(".env").exists():
    load_dotenv()

import anthropic
from rich.console import Console
from rich.panel import Panel
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

console = Console()

# ── Colores heru ──────────────────────────────────────────────────────────────
AZUL        = RGBColor(0x17, 0x90, 0xEC)   # #1790EC
AZUL_OSCURO = RGBColor(0x0C, 0x39, 0x61)   # #0C3961
BLANCO      = RGBColor(0xF4, 0xF9, 0xFE)   # #F4F9FE
NEGRO       = RGBColor(0x1A, 0x1A, 0x2E)

# ── Agente ────────────────────────────────────────────────────────────────────
def run(client, system, prompt, name, tokens=4000):
    console.print(f"\n[bold cyan]▶ {name}[/bold cyan]...")
    r = client.messages.create(
        model=os.getenv("MODEL", "claude-sonnet-4-6"),
        max_tokens=tokens,
        system=system,
        messages=[{"role": "user", "content": prompt}],
    )
    text = r.content[0].text.strip()
    console.print(Panel(text[:600] + ("…" if len(text) > 600 else ""),
                        title=f"[green]{name}[/green]", border_style="green"))
    return text

# ── Prompts ───────────────────────────────────────────────────────────────────

SYS_STRATEGY = """Eres el Director de Estrategia Digital de heru.app.
Conoces Google Ads, Meta Ads, TikTok Ads y YouTube Ads a nivel experto.
KPIs objetivo: CAC < $350 MXN | CPL < $150 MXN | ROAS > 4x
Nunca hagas preguntas. Entrega estrategia accionable con números reales."""

PROMPT_STRATEGY = """
BRIEF COMPLETO:
{brief}

PRESUPUESTO: $4,000 USD / $80,000 MXN mensual
Distribución propuesta: Google Search 40% | Meta 35% | TikTok 15% | YouTube 10%

NARRATIVA CENTRAL APROBADA:
"Marzo es para prepararte. En heru somos la ayuda que asegura tu tranquilidad fiscal ante el SAT,
no sólo durante un mes, sino por todo el año."

OFERTA: Plan anual = declaración anual GRATIS + 50% descuento en declaraciones atrasadas
Usar en 40% de piezas (conversión). 60% piezas = awareness/dolor.

Genera la estrategia completa en este formato JSON exacto (solo JSON, sin texto antes ni después):

{{
  "diagnostico": {{
    "propuesta_actual_bien": ["punto 1", "punto 2", "punto 3"],
    "propuesta_actual_mal": ["problema 1", "problema 2", "problema 3", "problema 4"],
    "errores_criticos": ["error 1", "error 2"]
  }},
  "objetivo_campana": "texto de 2-3 oraciones explicando el objetivo real (suscripción, no declaración)",
  "audiencias": [
    {{"segmento": "nombre", "descripcion": "quiénes son", "dolor_principal": "su miedo fiscal", "canal_prioritario": "canal"}}
  ],
  "canales": [
    {{
      "canal": "Google Ads Search",
      "por_que": "razón estratégica de 2 oraciones",
      "presupuesto_mxn": 32000,
      "presupuesto_usd": 1600,
      "objetivo": "objetivo específico",
      "kpi_principal": "CPL < $150 MXN",
      "grupos": [
        {{
          "nombre": "nombre del grupo",
          "intencion": "intención de búsqueda",
          "keywords": ["kw1", "kw2", "kw3", "kw4", "kw5"],
          "match_type": "phrase/exact/broad",
          "keywords_negativas": ["neg1", "neg2", "neg3"],
          "presupuesto_dia_mxn": 1000,
          "angulo": "awareness o conversion"
        }}
      ]
    }},
    {{
      "canal": "Meta Ads",
      "por_que": "razón estratégica",
      "presupuesto_mxn": 28000,
      "presupuesto_usd": 1400,
      "objetivo": "objetivo",
      "kpi_principal": "CPL < $120 MXN",
      "campanas": [
        {{
          "nombre": "nombre campaña",
          "objetivo_meta": "AWARENESS/TRAFFIC/LEADS",
          "formato": "Feed estático/Reel/Carrusel/Story",
          "audiencia": "descripción detallada de audiencia",
          "segmentacion": "intereses y comportamientos específicos",
          "presupuesto_dia_mxn": 600,
          "angulo": "awareness o conversion"
        }}
      ]
    }},
    {{
      "canal": "TikTok Ads",
      "por_que": "razón estratégica",
      "presupuesto_mxn": 12000,
      "presupuesto_usd": 600,
      "objetivo": "objetivo",
      "kpi_principal": "CPV < $0.5 MXN",
      "campanas": [
        {{
          "nombre": "nombre",
          "formato": "In-Feed Video",
          "duracion": "15-30 seg",
          "audiencia": "descripción",
          "presupuesto_dia_mxn": 400,
          "angulo": "awareness o humor"
        }}
      ]
    }},
    {{
      "canal": "YouTube Ads",
      "por_que": "razón estratégica",
      "presupuesto_mxn": 8000,
      "presupuesto_usd": 400,
      "objetivo": "objetivo",
      "kpi_principal": "CPV < $0.8 MXN",
      "campanas": [
        {{
          "nombre": "nombre",
          "tipo": "In-Stream skippable/Bumper 6seg",
          "audiencia": "descripción",
          "targeting": "intereses/palabras clave/temas",
          "presupuesto_dia_mxn": 267,
          "angulo": "awareness o conversion"
        }}
      ]
    }}
  ],
  "calendario": [
    {{"semana": 1, "acciones": ["acción 1", "acción 2"]}},
    {{"semana": 2, "acciones": ["acción 1", "acción 2"]}},
    {{"semana": 3, "acciones": ["acción 1", "acción 2"]}},
    {{"semana": 4, "acciones": ["acción 1", "acción 2"]}}
  ],
  "kpis_proyectados": {{
    "google_ads": {{"cpl_objetivo": "$150 MXN", "leads_proyectados": 213, "ctr_esperado": "4-6%"}},
    "meta_ads":   {{"cpl_objetivo": "$120 MXN", "leads_proyectados": 233, "ctr_esperado": "2-3%"}},
    "tiktok_ads": {{"cpv_objetivo": "$0.5 MXN", "vistas_proyectadas": 24000}},
    "youtube_ads":{{"cpv_objetivo": "$0.8 MXN", "vistas_proyectadas": 10000}}
  }}
}}
"""

SYS_COPY = """Eres el copywriter senior de heru.app.
heru SIEMPRE en minúsculas. Frases cortas. Verbos en presente. Sin tecnicismos.
El héroe es el trabajador independiente, nunca heru.
CRÍTICO: Verifica que CADA headline de Google sea ≤30 caracteres contando espacios.
CRÍTICO: No uses animaciones en diseño — solo fotos reales, objetos, íconos.
Nunca hagas preguntas."""

PROMPT_COPY = """
BRIEF:
{brief}

ESTRATEGIA APROBADA:
{strategy_json}

NARRATIVA: "Marzo es para prepararte. En heru somos la ayuda que asegura tu tranquilidad fiscal,
no sólo durante un mes, sino por todo el año."

OFERTA (usar en piezas de conversión): Plan anual = declaración anual GRATIS + 50% en atrasadas.

Genera el copy completo en JSON exacto (solo JSON):

{{
  "google_ads": [
    {{
      "grupo": "nombre del grupo",
      "angulo": "awareness o conversion",
      "headlines": [
        {{"texto": "máx 30 chars", "chars": 22}},
        {{"texto": "máx 30 chars", "chars": 25}},
        {{"texto": "máx 30 chars", "chars": 18}},
        {{"texto": "máx 30 chars", "chars": 28}},
        {{"texto": "máx 30 chars", "chars": 20}}
      ],
      "descriptions": [
        {{"texto": "máx 90 chars — puede incluir la oferta si es conversion", "chars": 88}},
        {{"texto": "máx 90 chars", "chars": 85}},
        {{"texto": "máx 90 chars", "chars": 90}}
      ],
      "sitelinks": [
        {{"texto": "texto corto", "descripcion": "descripción"}},
        {{"texto": "texto corto", "descripcion": "descripción"}}
      ],
      "callouts": ["callout 1", "callout 2", "callout 3"],
      "cta": "texto del botón"
    }}
  ],
  "meta_ads": [
    {{
      "anuncio": "nombre",
      "angulo": "awareness o conversion",
      "formato": "Feed estático/Reel/Carrusel/Story",
      "headline": "máx 40 chars",
      "primary_text": "texto principal máx 125 chars — conversacional, sin tecnicismos",
      "descripcion": "máx 30 chars",
      "cta": "botón",
      "slides": ["slide 1 texto", "slide 2 texto", "slide 3 texto — puede tener oferta si conversion"]
    }}
  ],
  "tiktok_ads": [
    {{
      "video": "nombre",
      "angulo": "awareness o humor",
      "hook_0_3seg": "texto pantalla — gancho en 3 seg",
      "vo_desarrollo": "guión voz en off máx 60 palabras — tono casual",
      "cierre": "frase final + CTA — puede incluir oferta",
      "caption": "texto del post",
      "hashtags": ["#SAT", "#Freelancer", "#Heru", "#Impuestos", "#DeclaraciónAnual"]
    }}
  ],
  "youtube_ads": [
    {{
      "video": "nombre",
      "tipo": "In-Stream skippable 20-30seg o Bumper 6seg",
      "hook_inskippable": "primeros 5 seg — inskippable — gancho fuerte",
      "desarrollo": "guión máx 40 palabras (In-Stream) o 6 palabras (Bumper)",
      "cta_final": "pantalla final con URL"
    }}
  ]
}}
"""

SYS_DESIGN = """Eres el director creativo de heru.app.
REGLAS NO NEGOCIABLES:
- Fotos REALES únicamente. CERO animaciones, CERO ilustraciones, CERO GIFs.
- Elementos: computadoras, teléfonos, avisos del SAT, recibos, íconos fiscales. NO obligatorio mostrar personas.
- El logo 'heru' y el copy DEBEN ser prominentes y legibles.
- Paleta: #1790EC (azul heru), #0C3961 (azul oscuro), #F4F9FE (blanco).
- Tipografía: Lexend Bold para headlines, Helvetica para cuerpo.
- Para videos: storyboard de escenas, NO prompt de imagen fija.
Nunca hagas preguntas."""

PROMPT_DESIGN = """
COPY APROBADO:
{copy_json}

Para CADA pieza del copy genera el concepto visual en JSON exacto (solo JSON):

{{
  "google_ads": [
    {{
      "grupo": "nombre del grupo",
      "copy_en_pieza": "headline principal exacto que va en la imagen",
      "concepto": "idea visual en una línea",
      "elemento_principal": "objeto/escena real — sin personas si no es necesario",
      "composicion": "descripción de zonas: qué va arriba/centro/abajo, dónde va logo y texto",
      "logo_posicion": "esquina inferior derecha / superior izquierda",
      "texto_en_imagen": "Lexend Bold, color blanco, zona superior, alto contraste",
      "paleta": ["#1790EC", "#F4F9FE"],
      "formato": "1:1 Feed / 4:5 / 16:9",
      "prompt_ia": "detailed prompt in English for Midjourney/DALL-E — MUST include the exact Spanish copy text as text overlay in the image"
    }}
  ],
  "meta_ads": [
    {{
      "anuncio": "nombre",
      "copy_en_pieza": "copy exacto que aparece en la imagen/video",
      "concepto": "idea visual",
      "elemento_principal": "qué se ve en la foto/video",
      "composicion": "descripción de layout",
      "logo_posicion": "posición del logo",
      "formato": "1:1 / 4:5 / 9:16",
      "prompt_ia": "detailed English prompt — real photo, heru blue #1790EC, exact copy as text overlay",
      "storyboard": "para Reels: escena 1 → escena 2 → escena 3 (solo si es video)"
    }}
  ],
  "tiktok_ads": [
    {{
      "video": "nombre",
      "copy_en_pieza": "hook text exacto en pantalla",
      "concepto": "idea visual del video",
      "storyboard": [
        {{"seg": "0-3", "escena": "qué se ve en pantalla", "texto_pantalla": "texto exacto", "vo": "voz en off"}},
        {{"seg": "3-20", "escena": "descripción", "texto_pantalla": "texto", "vo": "voz en off"}},
        {{"seg": "20-30", "escena": "cierre con logo", "texto_pantalla": "CTA + heru", "vo": "cierre"}}
      ],
      "formato": "9:16 vertical"
    }}
  ],
  "youtube_ads": [
    {{
      "video": "nombre",
      "copy_en_pieza": "hook exacto en pantalla",
      "concepto": "idea del video",
      "storyboard": [
        {{"seg": "0-5", "escena": "inskippable — qué se ve", "texto": "texto en pantalla"}},
        {{"seg": "5-25", "escena": "desarrollo", "texto": "texto"}},
        {{"seg": "25-30", "escena": "pantalla final con logo heru + URL", "texto": "CTA"}}
      ],
      "formato": "16:9"
    }}
  ]
}}
"""

# ── PowerPoint ────────────────────────────────────────────────────────────────

def add_slide(prs, layout_idx=6):
    layout = prs.slide_layouts[layout_idx]
    return prs.slides.add_slide(layout)

def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def textbox(slide, text, l, t, w, h, size=18, bold=False, color=BLANCO, align=PP_ALIGN.LEFT, wrap=True):
    txBox = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return txBox

def bullet_box(slide, items, l, t, w, h, size=14, color=BLANCO):
    txBox = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(size)
        p.font.color.rgb = color

def divider(slide, l, t, w, color=AZUL):
    from pptx.util import Pt as PtU
    line = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()

def build_ppt(strategy: dict, copy: dict, design: dict, fecha: str) -> str:
    prs = Presentation()
    prs.slide_width  = Inches(13.33)
    prs.slide_height = Inches(7.5)

    # ── Slide 1: Portada ──────────────────────────────────────────────────────
    s = add_slide(prs)
    bg(s, AZUL_OSCURO)
    textbox(s, "heru", 0.5, 0.4, 4, 0.8, size=48, bold=True, color=AZUL)
    textbox(s, "Estrategia de Campaña Digital", 0.5, 1.3, 9, 0.6, size=28, bold=True)
    textbox(s, "Temporada Declaración Anual 2025 · Marzo–Abril 2026", 0.5, 2.0, 9, 0.5, size=16, color=RGBColor(0xA5, 0xD2, 0xFA))
    textbox(s, "Presupuesto: $4,000 USD / $80,000 MXN mensual", 0.5, 2.6, 9, 0.4, size=14)
    textbox(s, "Canales: Google Ads · Meta · TikTok · YouTube", 0.5, 3.0, 9, 0.4, size=14)
    divider(s, 0.5, 4.0, 8)
    textbox(s, '"Marzo es para prepararte. heru es la ayuda que asegura\ntu tranquilidad fiscal, no solo un mes, sino todo el año."',
            0.5, 4.2, 10, 1.0, size=13, color=RGBColor(0xA5, 0xD2, 0xFA))
    textbox(s, fecha[:10], 11, 7.0, 2, 0.4, size=10, color=RGBColor(0x88, 0xAA, 0xCC))

    # ── Slide 2: Objetivo estratégico ─────────────────────────────────────────
    s = add_slide(prs)
    bg(s, AZUL_OSCURO)
    divider(s, 0.5, 0.5, 2.5)
    textbox(s, "Objetivo Estratégico", 0.5, 0.6, 9, 0.6, size=26, bold=True)
    textbox(s, strategy.get("objetivo_campana", ""), 0.5, 1.4, 12, 1.2, size=16, color=RGBColor(0xA5, 0xD2, 0xFA))
    divider(s, 0.5, 2.8, 12)
    textbox(s, "La declaración anual es el gancho de entrada — la suscripción es el negocio.", 0.5, 3.0, 12, 0.5, size=14, bold=True, color=AZUL)
    textbox(s, "Oferta activadora:", 0.5, 3.7, 4, 0.4, size=13, bold=True)
    textbox(s, "Plan anual = Declaración anual GRATIS + 50% descuento en declaraciones atrasadas",
            0.5, 4.1, 12, 0.6, size=14, color=RGBColor(0xA5, 0xD2, 0xFA))
    textbox(s, "Distribución: 60% piezas awareness/dolor · 40% piezas conversión con oferta",
            0.5, 4.9, 12, 0.4, size=12)

    # ── Slide 3: Diagnóstico de la propuesta ──────────────────────────────────
    s = add_slide(prs)
    bg(s, AZUL_OSCURO)
    divider(s, 0.5, 0.5, 3)
    textbox(s, "Diagnóstico de la Propuesta", 0.5, 0.6, 9, 0.6, size=26, bold=True)
    diag = strategy.get("diagnostico", {})

    textbox(s, "✅ Lo que funciona", 0.5, 1.4, 5.5, 0.4, size=14, bold=True, color=RGBColor(0x00, 0xC4, 0x8C))
    bullet_box(s, diag.get("propuesta_actual_bien", [])[:3], 0.5, 1.9, 5.5, 1.8, size=12)

    textbox(s, "❌ Lo que falta o está mal", 7.0, 1.4, 5.8, 0.4, size=14, bold=True, color=RGBColor(0xFF, 0x6B, 0x6B))
    bullet_box(s, diag.get("propuesta_actual_mal", [])[:4], 7.0, 1.9, 5.8, 2.2, size=12)

    divider(s, 0.5, 4.3, 12)
    textbox(s, "Errores críticos a corregir antes de publicar:", 0.5, 4.5, 12, 0.4, size=13, bold=True, color=RGBColor(0xFF, 0x6B, 0x6B))
    bullet_box(s, diag.get("errores_criticos", []), 0.5, 4.9, 12, 1.0, size=12, color=RGBColor(0xFF, 0xAA, 0xAA))

    # ── Slide 4: Audiencias ───────────────────────────────────────────────────
    s = add_slide(prs)
    bg(s, AZUL_OSCURO)
    divider(s, 0.5, 0.5, 2)
    textbox(s, "Audiencias Objetivo", 0.5, 0.6, 9, 0.6, size=26, bold=True)
    audiencias = strategy.get("audiencias", [])
    cols = [0.5, 3.4, 6.3, 9.2, 12.1]
    for i, aud in enumerate(audiencias[:4]):
        x = cols[i] if i < len(cols) else 0.5
        # Caja de audiencia
        box = s.shapes.add_shape(1, Inches(x), Inches(1.4), Inches(2.7), Inches(5.5))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0x14, 0x3A, 0x6B)
        box.line.color.rgb = AZUL
        textbox(s, aud.get("segmento", ""), x+0.1, 1.5, 2.5, 0.5, size=13, bold=True, color=AZUL)
        textbox(s, aud.get("descripcion", ""), x+0.1, 2.1, 2.5, 0.8, size=11)
        textbox(s, "Dolor:", x+0.1, 3.1, 2.5, 0.3, size=11, bold=True, color=RGBColor(0xFF, 0xAA, 0x44))
        textbox(s, aud.get("dolor_principal", ""), x+0.1, 3.4, 2.5, 0.8, size=11, color=RGBColor(0xFF, 0xCC, 0x88))
        textbox(s, f"Canal: {aud.get('canal_prioritario', '')}", x+0.1, 4.4, 2.5, 0.3, size=11, color=AZUL)

    # ── Slide 5: Presupuesto y distribución ───────────────────────────────────
    s = add_slide(prs)
    bg(s, AZUL_OSCURO)
    divider(s, 0.5, 0.5, 3)
    textbox(s, "Distribución de Presupuesto", 0.5, 0.6, 9, 0.6, size=26, bold=True)
    textbox(s, "$4,000 USD / $80,000 MXN mensual", 0.5, 1.3, 8, 0.4, size=14, color=AZUL)

    canales_data = [
        ("Google Ads Search", "40%", "$1,600 USD", "$32,000 MXN", "Alta intención · captura demanda activa"),
        ("Meta Ads",          "35%", "$1,400 USD", "$28,000 MXN", "Alcance masivo · awareness + conversión"),
        ("TikTok Ads",        "15%", "$600 USD",   "$12,000 MXN", "Freelancers jóvenes · formato nativo"),
        ("YouTube Ads",       "10%", "$400 USD",   "$8,000 MXN",  "Pre-roll · awareness de marca"),
    ]
    headers = ["Canal", "%", "USD", "MXN", "Razón estratégica"]
    col_x   = [0.4, 3.2, 4.5, 5.9, 7.3]
    col_w   = [2.7, 1.1, 1.3, 1.3, 5.8]

    y_h = 2.0
    for j, h in enumerate(headers):
        textbox(s, h, col_x[j], y_h, col_w[j], 0.35, size=12, bold=True, color=AZUL)
    divider(s, 0.4, 2.4, 12.5)

    for i, row in enumerate(canales_data):
        y = 2.5 + i * 0.9
        box = s.shapes.add_shape(1, Inches(0.4), Inches(y), Inches(12.5), Inches(0.75))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0x14, 0x3A, 0x6B) if i % 2 == 0 else AZUL_OSCURO
        box.line.fill.background()
        for j, cell in enumerate(row):
            c = AZUL if j == 0 else BLANCO
            textbox(s, cell, col_x[j]+0.05, y+0.1, col_w[j]-0.1, 0.55, size=12, color=c)

    # ── Slides 6-9: Estrategia por canal ─────────────────────────────────────
    canal_icons = {"Google Ads Search": "🔍", "Meta Ads": "📘", "TikTok Ads": "🎵", "YouTube Ads": "▶️"}
    for canal_data in strategy.get("canales", []):
        s = add_slide(prs)
        bg(s, AZUL_OSCURO)
        nombre = canal_data.get("canal", "Canal")
        icon   = canal_icons.get(nombre, "📢")
        divider(s, 0.5, 0.5, 2)
        textbox(s, f"{icon}  {nombre}", 0.5, 0.6, 9, 0.6, size=24, bold=True, color=AZUL)
        textbox(s, canal_data.get("por_que", ""), 0.5, 1.3, 8.5, 0.7, size=13, color=RGBColor(0xA5, 0xD2, 0xFA))

        mxn = canal_data.get("presupuesto_mxn", 0)
        usd = canal_data.get("presupuesto_usd", 0)
        textbox(s, f"${mxn:,} MXN / ${usd:,} USD  ·  {canal_data.get('kpi_principal','')}",
                0.5, 2.1, 8, 0.4, size=13, color=AZUL)
        divider(s, 0.5, 2.6, 12)

        grupos = canal_data.get("grupos", canal_data.get("campanas", []))
        for i, g in enumerate(grupos[:4]):
            x = 0.5 + i * 3.2
            nombre_g = g.get("nombre", g.get("nombre", f"Grupo {i+1}"))
            box = s.shapes.add_shape(1, Inches(x), Inches(2.8), Inches(3.0), Inches(4.3))
            box.fill.solid()
            box.fill.fore_color.rgb = RGBColor(0x14, 0x3A, 0x6B)
            box.line.color.rgb = AZUL
            textbox(s, nombre_g, x+0.1, 2.9, 2.8, 0.45, size=12, bold=True, color=AZUL)

            # Keywords o audiencia
            kws = g.get("keywords", [])
            aud = g.get("audiencia", g.get("segmentacion", ""))
            if kws:
                textbox(s, "Keywords:", x+0.1, 3.45, 2.8, 0.3, size=10, bold=True, color=RGBColor(0xA5, 0xD2, 0xFA))
                bullet_box(s, kws[:4], x+0.1, 3.75, 2.8, 1.2, size=10)
            elif aud:
                textbox(s, "Audiencia:", x+0.1, 3.45, 2.8, 0.3, size=10, bold=True, color=RGBColor(0xA5, 0xD2, 0xFA))
                textbox(s, aud, x+0.1, 3.75, 2.8, 0.8, size=10)

            presup = g.get("presupuesto_dia_mxn", 0)
            angulo = g.get("angulo", "")
            textbox(s, f"${presup}/día  ·  {angulo}", x+0.1, 6.6, 2.8, 0.35, size=10, color=AZUL)

    # ── Slide: Calendario ─────────────────────────────────────────────────────
    s = add_slide(prs)
    bg(s, AZUL_OSCURO)
    divider(s, 0.5, 0.5, 2.5)
    textbox(s, "Calendario de Activación — Marzo/Abril 2026", 0.5, 0.6, 12, 0.6, size=24, bold=True)
    cal = strategy.get("calendario", [])
    semana_x = [0.4, 3.5, 6.6, 9.7]
    for i, sem in enumerate(cal[:4]):
        x = semana_x[i]
        box = s.shapes.add_shape(1, Inches(x), Inches(1.5), Inches(3.0), Inches(5.5))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0x14, 0x3A, 0x6B)
        box.line.color.rgb = AZUL
        textbox(s, f"Semana {sem.get('semana', i+1)}", x+0.1, 1.6, 2.8, 0.4, size=14, bold=True, color=AZUL)
        divider(s, x+0.1, 2.1, 2.7, color=AZUL)
        bullet_box(s, sem.get("acciones", []), x+0.1, 2.2, 2.8, 4.5, size=11)

    # ── Slide: KPIs proyectados ───────────────────────────────────────────────
    s = add_slide(prs)
    bg(s, AZUL_OSCURO)
    divider(s, 0.5, 0.5, 2.5)
    textbox(s, "KPIs Proyectados", 0.5, 0.6, 9, 0.6, size=26, bold=True)
    textbox(s, "Con presupuesto $80,000 MXN/mes", 0.5, 1.3, 8, 0.4, size=14, color=AZUL)

    kpis = strategy.get("kpis_proyectados", {})
    kpi_cards = [
        ("Google Ads", kpis.get("google_ads", {}), 0.4),
        ("Meta Ads",   kpis.get("meta_ads", {}),   3.5),
        ("TikTok",     kpis.get("tiktok_ads", {}), 6.6),
        ("YouTube",    kpis.get("youtube_ads", {}), 9.7),
    ]
    for nombre_k, data, x in kpi_cards:
        box = s.shapes.add_shape(1, Inches(x), Inches(2.0), Inches(3.0), Inches(5.0))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0x14, 0x3A, 0x6B)
        box.line.color.rgb = AZUL
        textbox(s, nombre_k, x+0.1, 2.1, 2.8, 0.4, size=14, bold=True, color=AZUL)
        divider(s, x+0.1, 2.6, 2.7)
        y = 2.7
        for k, v in data.items():
            label = k.replace("_", " ").title()
            textbox(s, f"{label}:", x+0.1, y, 2.8, 0.3, size=11, bold=True, color=RGBColor(0xA5, 0xD2, 0xFA))
            textbox(s, str(v), x+0.1, y+0.3, 2.8, 0.35, size=13, bold=True, color=BLANCO)
            y += 0.75

    # ── Slide final ───────────────────────────────────────────────────────────
    s = add_slide(prs)
    bg(s, AZUL)
    textbox(s, "heru", 5.0, 1.8, 3.5, 1.2, size=72, bold=True, color=BLANCO, align=PP_ALIGN.CENTER)
    textbox(s, "Tu acompañante fiscal todo el año.", 3.0, 3.2, 7.5, 0.6, size=18, align=PP_ALIGN.CENTER)
    textbox(s, "heru.app", 5.5, 4.2, 2.5, 0.5, size=16, align=PP_ALIGN.CENTER, color=AZUL_OSCURO)

    # ── Guardar ───────────────────────────────────────────────────────────────
    out = Path("output/reports")
    out.mkdir(parents=True, exist_ok=True)
    path = out / f"heru_campana_{fecha}.pptx"
    prs.save(str(path))
    return str(path)

# ── CSV builders ──────────────────────────────────────────────────────────────

def save_csv(rows, filename):
    if not rows:
        return None
    out = Path("output/reports")
    out.mkdir(parents=True, exist_ok=True)
    path = out / filename
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=rows[0].keys())
        w.writeheader()
        w.writerows(rows)
    return str(path)


def build_google_csv(copy: dict, design: dict) -> list[dict]:
    rows = []
    design_map = {d.get("grupo", ""): d for d in design.get("google_ads", [])}
    for g in copy.get("google_ads", []):
        grupo = g.get("grupo", "")
        d     = design_map.get(grupo, {})
        hs    = g.get("headlines", [])
        ds    = g.get("descriptions", [])
        sls   = g.get("sitelinks", [])
        rows.append({
            "Grupo":           grupo,
            "Ángulo":          g.get("angulo", ""),
            "H1 (≤30)":       hs[0].get("texto","") if len(hs) > 0 else "",
            "H2 (≤30)":       hs[1].get("texto","") if len(hs) > 1 else "",
            "H3 (≤30)":       hs[2].get("texto","") if len(hs) > 2 else "",
            "H4 (≤30)":       hs[3].get("texto","") if len(hs) > 3 else "",
            "H5 (≤30)":       hs[4].get("texto","") if len(hs) > 4 else "",
            "D1 (≤90)":       ds[0].get("texto","") if len(ds) > 0 else "",
            "D2 (≤90)":       ds[1].get("texto","") if len(ds) > 1 else "",
            "D3 (≤90)":       ds[2].get("texto","") if len(ds) > 2 else "",
            "Sitelink 1":      sls[0].get("texto","") if len(sls) > 0 else "",
            "Sitelink 2":      sls[1].get("texto","") if len(sls) > 1 else "",
            "Callouts":        " | ".join(g.get("callouts", [])),
            "CTA":             g.get("cta", ""),
            "Copy en pieza":   d.get("copy_en_pieza", ""),
            "Concepto visual": d.get("concepto", ""),
            "Elemento":        d.get("elemento_principal", ""),
            "Composición":     d.get("composicion", ""),
            "Prompt IA":       d.get("prompt_ia", ""),
            "Formato":         d.get("formato", ""),
            "Estado":          "Pendiente revisión",
        })
    return rows


def build_meta_csv(copy: dict, design: dict) -> list[dict]:
    rows = []
    design_map = {d.get("anuncio", ""): d for d in design.get("meta_ads", [])}
    for a in copy.get("meta_ads", []):
        nombre = a.get("anuncio", "")
        d      = design_map.get(nombre, {})
        slides = a.get("slides", [])
        rows.append({
            "Anuncio":         nombre,
            "Ángulo":          a.get("angulo", ""),
            "Formato":         a.get("formato", ""),
            "Headline":        a.get("headline", ""),
            "Primary Text":    a.get("primary_text", ""),
            "Descripción":     a.get("descripcion", ""),
            "CTA":             a.get("cta", ""),
            "Slide 1":         slides[0] if len(slides) > 0 else "",
            "Slide 2":         slides[1] if len(slides) > 1 else "",
            "Slide 3":         slides[2] if len(slides) > 2 else "",
            "Copy en pieza":   d.get("copy_en_pieza", ""),
            "Concepto visual": d.get("concepto", ""),
            "Elemento":        d.get("elemento_principal", ""),
            "Composición":     d.get("composicion", ""),
            "Prompt IA":       d.get("prompt_ia", ""),
            "Storyboard":      d.get("storyboard", ""),
            "Formato visual":  d.get("formato", ""),
            "Estado":          "Pendiente revisión",
        })
    return rows


def build_tiktok_csv(copy: dict, design: dict) -> list[dict]:
    rows = []
    design_map = {d.get("video", ""): d for d in design.get("tiktok_ads", [])}
    for v in copy.get("tiktok_ads", []):
        nombre = v.get("video", "")
        d      = design_map.get(nombre, {})
        sb     = d.get("storyboard", [])
        rows.append({
            "Video":           nombre,
            "Ángulo":          v.get("angulo", ""),
            "Hook 0-3seg":     v.get("hook_0_3seg", ""),
            "VO Desarrollo":   v.get("vo_desarrollo", ""),
            "Cierre + CTA":    v.get("cierre", ""),
            "Caption":         v.get("caption", ""),
            "Hashtags":        " ".join(v.get("hashtags", [])),
            "Copy en pieza":   d.get("copy_en_pieza", ""),
            "Concepto":        d.get("concepto", ""),
            "Escena 0-3seg":   sb[0].get("escena","") if len(sb) > 0 else "",
            "Escena 3-20seg":  sb[1].get("escena","") if len(sb) > 1 else "",
            "Escena 20-30seg": sb[2].get("escena","") if len(sb) > 2 else "",
            "Formato":         "9:16 vertical",
            "Estado":          "Pendiente revisión",
        })
    return rows


def build_youtube_csv(copy: dict, design: dict) -> list[dict]:
    rows = []
    design_map = {d.get("video", ""): d for d in design.get("youtube_ads", [])}
    for v in copy.get("youtube_ads", []):
        nombre = v.get("video", "")
        d      = design_map.get(nombre, {})
        sb     = d.get("storyboard", [])
        rows.append({
            "Video":            nombre,
            "Tipo":             v.get("tipo", ""),
            "Hook inskippable": v.get("hook_inskippable", ""),
            "Desarrollo":       v.get("desarrollo", ""),
            "CTA Final":        v.get("cta_final", ""),
            "Copy en pieza":    d.get("copy_en_pieza", ""),
            "Concepto":         d.get("concepto", ""),
            "Escena 0-5seg":    sb[0].get("escena","") if len(sb) > 0 else "",
            "Escena 5-25seg":   sb[1].get("escena","") if len(sb) > 1 else "",
            "Escena final":     sb[2].get("escena","") if len(sb) > 2 else "",
            "Formato":          "16:9",
            "Estado":           "Pendiente revisión",
        })
    return rows

# ── Parse JSON del agente ─────────────────────────────────────────────────────

def parse_json(text: str) -> dict:
    """Extrae y parsea el JSON del texto del agente.
    Soporta: JSON puro, bloque ```json ... ```, o JSON embebido en texto.
    """
    # 1. Intentar extraer de bloque de código ```json ... ```
    code_block = re.search(r'```(?:json)?\s*(\{[\s\S]*?\})\s*```', text)
    if code_block:
        try:
            return json.loads(code_block.group(1))
        except json.JSONDecodeError:
            pass

    # 2. Buscar el JSON más grande (de { hasta la última })
    match = re.search(r'\{[\s\S]*\}', text)
    if not match:
        console.print("[red]No se encontró ningún bloque JSON en la respuesta.[/red]")
        console.print(f"[dim]Primeros 500 chars: {text[:500]}[/dim]")
        return {}
    try:
        return json.loads(match.group())
    except json.JSONDecodeError as e:
        console.print(f"[yellow]JSON parse error: {e}[/yellow]")
        # 3. Reparar JSON truncado añadiendo llaves de cierre
        raw = match.group()
        for extra in range(1, 15):
            try:
                return json.loads(raw + "}" * extra)
            except Exception:
                pass
        console.print("[red]No se pudo reparar el JSON. Muestra del texto:[/red]")
        console.print(f"[dim]{text[-400:]}[/dim]")
        return {}

# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    api_key = os.getenv("ANTHROPIC_API_KEY")
    if not api_key:
        console.print("[red]❌ ANTHROPIC_API_KEY no encontrada en .env[/red]")
        sys.exit(1)

    client = anthropic.Anthropic(api_key=api_key)
    fecha  = datetime.now().strftime("%Y%m%d_%H%M")

    console.print(Panel(
        "[bold green]Constructor de Campañas — heru.app[/bold green]\n"
        "Google Ads · Meta · TikTok · YouTube\n"
        "Entregables: PPT estratégico + 4 CSVs por canal",
        border_style="green",
    ))

    # ── Brief ─────────────────────────────────────────────────────────────────
    default_brief = Path(__file__).parent.parent / "output" / "brief_tofu.txt"
    if default_brief.exists():
        brief = default_brief.read_text(encoding="utf-8")
        console.print(f"[dim]Brief: {default_brief}[/dim]")
    else:
        console.print("[yellow]Pega el brief y presiona Enter dos veces.[/yellow]")
        lines = []
        while True:
            line = input()
            if line == "" and lines and lines[-1] == "":
                break
            lines.append(line)
        brief = "\n".join(lines).strip()

    # ── Agente 1: Estrategia ──────────────────────────────────────────────────
    strategy_raw = run(client, SYS_STRATEGY,
                       PROMPT_STRATEGY.format(brief=brief),
                       "Performance Ads — Estrategia", tokens=8000)
    strategy = parse_json(strategy_raw)
    if not strategy:
        console.print("[red]⚠ No pude parsear la estrategia. Revisa el output arriba.[/red]")
        sys.exit(1)

    # ── Agente 2: Copy ────────────────────────────────────────────────────────
    copy_raw = run(client, SYS_COPY,
                   PROMPT_COPY.format(brief=brief, strategy_json=json.dumps(strategy, ensure_ascii=False, indent=2)),
                   "Copywriter", tokens=8000)
    copy_data = parse_json(copy_raw)

    # ── Agente 3: Diseño ──────────────────────────────────────────────────────
    design_raw = run(client, SYS_DESIGN,
                     PROMPT_DESIGN.format(copy_json=json.dumps(copy_data, ensure_ascii=False, indent=2)),
                     "Graphic Designer", tokens=5000)
    design_data = parse_json(design_raw)

    # ── Generar archivos ──────────────────────────────────────────────────────
    console.print("\n[bold]Generando entregables...[/bold]")

    ppt_path = build_ppt(strategy, copy_data, design_data, fecha)

    files = [("PPT Estrategia", ppt_path)]

    google_rows = build_google_csv(copy_data, design_data)
    meta_rows   = build_meta_csv(copy_data, design_data)
    tiktok_rows = build_tiktok_csv(copy_data, design_data)
    yt_rows     = build_youtube_csv(copy_data, design_data)

    if google_rows:
        p = save_csv(google_rows, f"google_ads_{fecha}.csv")
        files.append(("Google Ads CSV", p))
    if meta_rows:
        p = save_csv(meta_rows, f"meta_ads_{fecha}.csv")
        files.append(("Meta Ads CSV", p))
    if tiktok_rows:
        p = save_csv(tiktok_rows, f"tiktok_ads_{fecha}.csv")
        files.append(("TikTok CSV", p))
    if yt_rows:
        p = save_csv(yt_rows, f"youtube_ads_{fecha}.csv")
        files.append(("YouTube CSV", p))

    files_str = "\n".join(f"  • {n}: [cyan]{p}[/cyan]" for n, p in files)
    console.print(Panel(
        f"[bold green]✅ Listo — {len(files)} archivos generados[/bold green]\n\n"
        f"{files_str}\n\n"
        "[dim]PPT: ábrelo en PowerPoint o Google Slides (Archivo → Importar)[/dim]\n"
        "[dim]CSVs: Google Sheets → Archivo → Importar → una pestaña por canal[/dim]",
        border_style="green",
    ))


if __name__ == "__main__":
    main()
